"""Выгрузка в Excel: сводка, портфель (по аккаунтам), закрытые позиции, ордера, журнал.

Деньги — числами с форматом валюты, проценты — долями с форматом «%», даты — датами Excel,
чтобы в Excel можно было сразу сортировать, фильтровать и считать.
"""

import io
import sqlite3
from collections.abc import Iterable
from datetime import UTC, date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from steamfolio import stats, views
from steamfolio.clock import MSK
from steamfolio.ledger import Ledger, Op

MONEY = {"USD": '"$"#,##0.00', "RUB": '#,##0.00 "₽"', "CNY": '#,##0.00 "¥"', "UAH": '#,##0.00 "₴"'}
PRICE = {"USD": '"$"#,##0.00##', "RUB": '#,##0.00## "₽"', "CNY": '#,##0.00## "¥"', "UAH": '#,##0.00## "₴"'}
PCT = "0.00%"
DATE = "DD.MM.YYYY"
DATETIME = "DD.MM.YYYY HH:MM"
HEADER_FILL = PatternFill("solid", fgColor="DCE6F2")

# Ключ строки портфеля: (предмет, id аккаунта, валюта).
RowKey = tuple[str, int, str]


def _day(iso: str | None) -> date | None:
    return date.fromisoformat(iso) if iso else None


def _msk(ts: str | None) -> datetime | None:
    """Excel не знает часовых поясов — пишем московское время без пояса."""
    return datetime.fromisoformat(ts).astimezone(MSK).replace(tzinfo=None) if ts else None


def _frac(pct: float | None) -> float | None:
    return None if pct is None else pct / 100


class _Sheet:
    """Лист-таблица: заголовок, строки, форматы колонок, автофильтр и закреплённая шапка."""

    def __init__(self, ws: Worksheet, columns: list[tuple[str, int]]):
        self.ws = ws
        ws.append([title for title, _ in columns])
        for i, (_, width) in enumerate(columns, 1):
            cell = ws.cell(1, i)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "B2"

    def add(self, values: list[Any], formats: dict[int, str] | None = None, link: tuple[int, str] | None = None) -> None:
        self.ws.append(values)
        row = self.ws.max_row
        for col, fmt in (formats or {}).items():
            self.ws.cell(row, col).number_format = fmt
        if link:
            cell = self.ws.cell(row, link[0])
            cell.hyperlink = link[1]
            cell.font = Font(color="0563C1", underline="single")

    def finish(self) -> None:
        if self.ws.max_row > 1:
            self.ws.auto_filter.ref = self.ws.dimensions


def _positions(sheet: _Sheet, rows: Iterable[dict[str, Any]]) -> None:
    for r in rows:
        cur = r["currency"]
        sheet.add(
            [r["item"], r["type"], r["account"], cur, r["quantity"], r["avg_price"], r["avg_price_usd"],
             _day(r["first_buy_date"]), r["price"], _msk(r["price_at"]), _frac(r["pct_avg"]), _frac(r["pct_1d"]),
             _frac(r["pct_7d"]), _frac(r["pct_30d"]), r.get("value_usd"), r["cost_usd"], r["comment"]],
            {6: PRICE[cur], 7: PRICE["USD"], 8: DATE, 9: MONEY[cur], 10: DATETIME, 11: PCT, 12: PCT, 13: PCT,
             14: PCT, 15: MONEY["USD"], 16: MONEY["USD"]},
            link=(1, r["url"]),
        )


POSITION_COLUMNS = [
    ("Название", 40), ("Тип", 18), ("Аккаунт", 16), ("Валюта", 8), ("Кол-во", 9), ("Средняя цена", 14),
    ("Средняя цена, USD", 14), ("Первая покупка", 13), ("Текущая цена (лучший ордер)", 14), ("Цена снята, МСК", 17),
    ("% к средней", 11), ("% сутки", 10), ("% неделя", 10), ("% месяц", 10), ("Стоимость, USD", 14),
    ("Себестоимость, USD", 14), ("Комментарий", 30),
]


def _status(r: dict[str, Any]) -> str:
    parts = ["откачена" if r["cancelled"] else "действует"]
    if r["edited_by"]:
        parts.append("изменена: " + ", ".join(f"#{i}" for i in r["edited_by"]))
    if r["deleted_by"]:
        parts.append("удалена: " + ", ".join(f"#{i}" for i in r["deleted_by"]))
    return "; ".join(parts)


def _journal(sheet: _Sheet, rows: list[dict[str, Any]]) -> None:
    for r in sorted(rows, key=lambda r: r["id"]):
        cur = r["currency"]
        total = round(r["price"] * r["quantity"], 2) if r["price"] is not None and r["quantity"] else None
        relation = []
        if r["target_id"]:
            relation.append(f"к операции #{r['target_id']} ({r['target_kind']})")
        if r["order_id"]:
            relation.append(f"ордер #{r['order_id']}")
        changes = "; ".join(f"{c['field']}: {c['before']} → {c['after']}" for c in r["changes"])
        source = {"usd": "", "steam": "курс Steam", "cbr": "курс ЦБ на дату", "manual": "вручную"}.get(r["fx_source"] or "", "")
        sheet.add(
            [r["id"], _day(r["op_date"]), _msk(r["created_at"]), r["kind_name"], r["item"], r["account"], r["to_account"],
             r["quantity"], r["price"], cur, total, r["fx_rate"], source, r["sum_usd"], _status(r),
             ", ".join(relation), changes, r["comment"]],
            {2: DATE, 3: DATETIME, 9: MONEY.get(cur, "0.00"), 11: MONEY.get(cur, "0.00"), 12: "0.0000",
             14: MONEY["USD"]},
        )


def build(
    conn: sqlite3.Connection, ledger: Ledger, ops: list[Op], portfolio_id: int = 1,
    only: dict[str, set[RowKey]] | None = None, now: datetime | None = None,
) -> bytes:
    """only — строки портфеля, оставшиеся после фильтров на сайте: {"open": {...}, "closed": {...}}; None — всё."""
    now = now or datetime.now(UTC)
    wb = Workbook()

    def keep(rows: list[dict[str, Any]], part: str) -> list[dict[str, Any]]:
        if only is None:
            return rows
        return [r for r in rows if (r["item"], r["account_id"], r["currency"]) in only.get(part, set())]

    open_rows = keep(views.portfolio(conn, ledger, portfolio_id=portfolio_id, now=now)["rows"], "open")
    closed_rows = keep(views.portfolio(conn, ledger, closed=True, portfolio_id=portfolio_id, now=now)["rows"], "closed")

    # Сводка
    ws = wb.active
    ws.title = "Сводка"
    s = stats.summary(conn, ledger, portfolio_id)
    lines: list[tuple[str, Any, str | None]] = [
        ("Портфолио Steam — выгрузка", None, None),
        ("Выгружено, МСК", _msk(now.isoformat()), DATETIME),
        ("Цены обновлены, МСК", _msk(views.last_update(conn)), DATETIME),
        ("Фильтры с сайта", "учтены (листы «Портфель» и «Закрытые позиции»)" if only is not None else "не учтены — выгружено всё", None),
        ("", None, None),
        ("Вложено в текущие позиции", s["invested_usd"], MONEY["USD"]),
        ("Текущая стоимость", s["value_usd"], MONEY["USD"]),
        ("Нереализованная прибыль", s["profit_usd"], MONEY["USD"]),
        ("Всего вложено", s["total_invested_usd"], MONEY["USD"]),
        ("Всего продано", s["total_sold_usd"], MONEY["USD"]),
        ("Всего получено прибыли", s["total_profit_usd"], MONEY["USD"]),
    ]
    if s["unpriced"]:
        lines.append(("Позиций без цены (учтены по себестоимости)", s["unpriced"], None))
    for label, value, fmt in lines:
        ws.append([label, value])
        if fmt:
            ws.cell(ws.max_row, 2).number_format = fmt
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 48
    ws.cell(6, 1).font = Font(bold=True)

    sheet = _Sheet(wb.create_sheet("Портфель"), POSITION_COLUMNS)
    _positions(sheet, open_rows)
    sheet.finish()

    sheet = _Sheet(wb.create_sheet("Закрытые позиции"), POSITION_COLUMNS)
    _positions(sheet, closed_rows)
    sheet.finish()

    orders = views.orders(conn, ledger, portfolio_id)
    rates = {"USD": 1.0} | {c.name: r for c, r in views.latest_rates(conn).items()}
    sheet = _Sheet(wb.create_sheet("Ордера"), [
        ("№ ордера", 10), ("Название", 40), ("Аккаунт", 16), ("Валюта", 8), ("Кол-во", 9), ("Цена ордера", 13),
        ("Сумма", 14), ("Сумма в USD по текущему курсу Steam", 16), ("Лучший ордер на площадке", 14),
        ("Перебит", 9), ("Комментарий", 30),
    ])
    for r in orders["rows"]:
        cur = r["currency"]
        usd = round(r["total"] / rates[cur], 2) if cur in rates else None
        sheet.add(
            [r["order_id"], r["item"], r["account"], cur, r["quantity"], r["price"], r["total"], usd, r["market_best"],
             "да" if r["outbid"] else "", r["comment"]],
            {6: MONEY[cur], 7: MONEY[cur], 8: MONEY["USD"], 9: MONEY[cur]},
            link=(2, r["url"]),
        )
    sheet.finish()

    sheet = _Sheet(wb.create_sheet("Журнал"), [
        ("№", 7), ("Дата операции", 13), ("Записана, МСК", 17), ("Тип", 22), ("Предмет", 40), ("Аккаунт", 16),
        ("На аккаунт", 16), ("Кол-во", 9), ("Цена", 13), ("Валюта", 8), ("Сумма", 14), ("Курс к USD", 11),
        ("Источник курса", 16), ("Сумма в USD по курсу на дату", 15), ("Статус", 24), ("Связь", 28),
        ("Изменения", 40), ("Комментарий", 30),
    ])
    _journal(sheet, views.journal(conn, ops, portfolio_id)["rows"])
    sheet.finish()

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
