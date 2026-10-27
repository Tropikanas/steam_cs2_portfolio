import io
from datetime import UTC, date, datetime

from openpyxl import load_workbook

from steamfolio import export
from steamfolio.currencies import Currency
from steamfolio.ledger import replay

RUB, USD = Currency.RUB, Currency.USD
CASE, AK, AWP = "Kilowatt Case", "AK-47 | Redline (Field-Tested)", "AWP | Asiimov (Field-Tested)"
D = date.fromisoformat


def rows(ws):
    header = [c.value for c in ws[1]]
    return [dict(zip(header, (c.value for c in r))) for r in ws.iter_rows(min_row=2)]


def setup(conn, journal, acc, acc2):
    journal.buy(CASE, acc, RUB, 10, 800, D("2026-09-28"))  # $1.00 по курсу 80
    journal.transfer(CASE, acc, acc2, RUB, 4, D("2026-09-29"))
    journal.buy(AK, acc, USD, 1, 1000, D("2026-09-28"))
    journal.sell(AK, acc, USD, 1, 1500, D("2026-10-01"))
    bad = journal.buy(AWP, acc, USD, 1, 3000, D("2026-10-02"))
    journal.cancel(bad)
    order = journal.place_order(CASE, acc, RUB, 5, 700)
    journal.fill_order(order, 2, D("2026-10-03"))
    with conn:
        run = conn.execute(
            "INSERT INTO collector_runs (started_at, finished_at, status) VALUES (?, ?, 'ok')",
            ("2026-10-27T18:00:00+00:00",) * 2,
        ).lastrowid
        conn.execute(
            """INSERT INTO price_snapshots (run_id, taken_at, market_hash_name, currency, best_buy, lowest_sell, source)
               VALUES (?, '2026-10-27T18:00:00+00:00', ?, ?, 900, NULL, 'histogram')""", (run, CASE, int(RUB)))


def workbook(conn, journal, only=None):
    ops = journal.ops()
    data = export.build(conn, replay(ops), ops, only=only, now=datetime(2026, 10, 27, 19, 0, tzinfo=UTC))
    return load_workbook(io.BytesIO(data))


def test_full_export(conn, journal, acc, acc2):
    setup(conn, journal, acc, acc2)
    wb = workbook(conn, journal)
    assert wb.sheetnames == ["Сводка", "Портфель", "Закрытые позиции", "Ордера", "Журнал"]

    portfolio = rows(wb["Портфель"])
    by_account = {(r["Название"], r["Аккаунт"]): r for r in portfolio}
    assert {k: r["Кол-во"] for k, r in by_account.items()} == {(CASE, "Основной"): 8, (CASE, "Второй"): 4}
    main = by_account[(CASE, "Основной")]
    assert main["Текущая цена (лучший ордер)"] == 9.0
    assert main["Первая покупка"] == datetime(2026, 9, 28)
    assert wb["Портфель"]["F2"].number_format.endswith('"₽"')
    assert wb["Портфель"]["A2"].hyperlink.target.startswith("https://steamcommunity.com/market/listings/730/")

    assert [r["Название"] for r in rows(wb["Закрытые позиции"])] == [AK]
    orders = rows(wb["Ордера"])
    assert len(orders) == 1 and orders[0]["Кол-во"] == 3 and orders[0]["Сумма"] == 21.0

    journal_rows = rows(wb["Журнал"])
    assert [r["№"] for r in journal_rows] == sorted(r["№"] for r in journal_rows)
    sale = next(r for r in journal_rows if r["Тип"] == "продажа")
    assert sale["Сумма в USD по курсу на дату"] == 15.0
    cancelled = next(r for r in journal_rows if r["Предмет"] == AWP)
    assert cancelled["Статус"].startswith("откачена")
    transfer = next(r for r in journal_rows if r["Тип"] == "перенос на аккаунт")
    assert (transfer["Аккаунт"], transfer["На аккаунт"], transfer["Цена"]) == ("Основной", "Второй", None)

    summary = {r[0].value: r[1].value for r in wb["Сводка"].iter_rows()}
    assert summary["Фильтры с сайта"].startswith("не учтены")
    assert summary["Всего продано"] == 15.0


def test_export_with_site_filters(conn, journal, acc, acc2):
    setup(conn, journal, acc, acc2)
    wb = workbook(conn, journal, only={"open": {(CASE, acc2, "RUB")}, "closed": set()})
    assert [(r["Название"], r["Аккаунт"]) for r in rows(wb["Портфель"])] == [(CASE, "Второй")]
    assert rows(wb["Закрытые позиции"]) == []
    assert len(rows(wb["Журнал"])) == len(journal.ops())  # журнал — всегда целиком
